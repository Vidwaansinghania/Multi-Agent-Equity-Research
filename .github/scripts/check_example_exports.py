"""Check the worked example's exports have the shape CLAUDE.md promises.

Usage: python .github/scripts/check_example_exports.py <out_dir>

<out_dir> holds workpapers.docx, model.xlsx and report.pdf plus the builders'
stdout as workpapers.log, workbook.log and report.log.
"""

import os
import re
import sys

import docx
import openpyxl

SECTIONS = 12
SHEETS = 14
PAGES = 6


def pdf_pages(path):
    # reportlab writes the page tree root uncompressed, so its /Count is readable as bytes.
    with open(path, "rb") as f:
        data = f.read()
    counts = re.findall(rb"/Type\s*/Pages\b[^>]*?/Count\s+(\d+)", data)
    counts += re.findall(rb"/Count\s+(\d+)[^>]*?/Type\s*/Pages\b", data)
    if not counts:
        sys.exit("could not find the page tree in %s" % path)
    return max(int(c) for c in counts)


def main():
    out = sys.argv[1]
    failures = []

    def check(label, got, want):
        ok = got == want
        print("%-4s %-22s %s (expected %s)" % ("ok" if ok else "FAIL", label, got, want))
        if not ok:
            failures.append(label)

    log = open(os.path.join(out, "workpapers.log"), encoding="utf-8").read()
    m = re.search(r"^sections: (\d+)$", log, re.M)
    check("DOCX sections", int(m.group(1)) if m else None, SECTIONS)
    absent = re.search(r"^absent: (.*)$", log, re.M)
    check("DOCX absent sources", absent.group(1) if absent else "none", "none")
    docx.Document(os.path.join(out, "workpapers.docx"))  # raises if the file is unreadable

    wb = openpyxl.load_workbook(os.path.join(out, "model.xlsx"), read_only=True)
    check("XLSX sheets", len(wb.sheetnames), SHEETS)

    check("PDF pages", pdf_pages(os.path.join(out, "report.pdf")), PAGES)

    if failures:
        sys.exit("failed: %s" % ", ".join(failures))


if __name__ == "__main__":
    main()
